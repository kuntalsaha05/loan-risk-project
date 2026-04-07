from pathlib import Path

from pptx import Presentation
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE
from pptx.util import Inches, Pt


# 1. Set file paths.
base_dir = Path(__file__).resolve().parent
visuals_dir = base_dir / "visuals"
output_file = base_dir / "Loan_Risk_Project_Presentation.pptx"


# 2. Create a new presentation.
prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)


# 3. Add title slide.
slide = prs.slides.add_slide(prs.slide_layouts[0])
slide.shapes.title.text = "Loan Risk Analysis and Prediction System"
slide.placeholders[1].text = "DET Mini Project Presentation"


# 4. Add project objective slide.
slide = prs.slides.add_slide(prs.slide_layouts[1])
slide.shapes.title.text = "Project Objective"
text_frame = slide.placeholders[1].text_frame
text_frame.clear()
bullets = [
    "Analyze a real loan dataset and identify default-related patterns",
    "Apply preprocessing, ETL, association mining, classification, and clustering",
    "Present insights using charts, reports, and a prediction module",
]
for index, bullet in enumerate(bullets):
    if index == 0:
        paragraph = text_frame.paragraphs[0]
    else:
        paragraph = text_frame.add_paragraph()
    paragraph.text = bullet
    paragraph.font.size = Pt(22)


# 5. Add dataset exploration slide.
slide = prs.slides.add_slide(prs.slide_layouts[1])
slide.shapes.title.text = "Dataset Exploration"
text_frame = slide.placeholders[1].text_frame
text_frame.clear()
bullets = [
    "Dataset: lending_club_loan_two.csv",
    "Rows: 396,030 and selected features: 24",
    "Explored numerical, categorical, and time-related fields",
    "Key risk indicators: loan amount, income, DTI, loan status, grade",
]
for index, bullet in enumerate(bullets):
    if index == 0:
        paragraph = text_frame.paragraphs[0]
    else:
        paragraph = text_frame.add_paragraph()
    paragraph.text = bullet
    paragraph.font.size = Pt(22)


# 6. Add preprocessing slide.
slide = prs.slides.add_slide(prs.slide_layouts[1])
slide.shapes.title.text = "Data Preprocessing"
text_frame = slide.placeholders[1].text_frame
text_frame.clear()
bullets = [
    "Handled missing values and removed unnecessary columns",
    "Clipped outliers using IQR-based limits",
    "Transformed date fields into issue year, issue month, and credit history years",
    "Created engineered features: default_flag and default_stage",
]
for index, bullet in enumerate(bullets):
    if index == 0:
        paragraph = text_frame.paragraphs[0]
    else:
        paragraph = text_frame.add_paragraph()
    paragraph.text = bullet
    paragraph.font.size = Pt(22)


# 7. Add algorithms slide.
slide = prs.slides.add_slide(prs.slide_layouts[1])
slide.shapes.title.text = "Algorithms Used"
text_frame = slide.placeholders[1].text_frame
text_frame.clear()
bullets = [
    "Apriori for association rule mining",
    "Logistic Regression, Decision Tree, and Random Forest for classification",
    "K-Means clustering for low, medium, and high risk groups",
    "Matplotlib and Seaborn for visualization",
]
for index, bullet in enumerate(bullets):
    if index == 0:
        paragraph = text_frame.paragraphs[0]
    else:
        paragraph = text_frame.add_paragraph()
    paragraph.text = bullet
    paragraph.font.size = Pt(22)


# 8. Add supervised learning results table.
slide = prs.slides.add_slide(prs.slide_layouts[5])
slide.shapes.title.text = "Supervised Learning Results"
table = slide.shapes.add_table(
    rows=4,
    cols=3,
    left=Inches(0.5),
    top=Inches(1.5),
    width=Inches(9),
    height=Inches(4.0),
).table

table.cell(0, 0).text = "Model"
table.cell(0, 1).text = "Accuracy"
table.cell(0, 2).text = "F1 Score"

table.cell(1, 0).text = "Logistic Regression"
table.cell(1, 1).text = "0.8056"
table.cell(1, 2).text = "0.0876"

table.cell(2, 0).text = "Decision Tree"
table.cell(2, 1).text = "0.8017"
table.cell(2, 2).text = "0.1476"

table.cell(3, 0).text = "Random Forest"
table.cell(3, 1).text = "0.8047"
table.cell(3, 2).text = "0.0358"


# 9. Add clustering results table.
slide = prs.slides.add_slide(prs.slide_layouts[5])
slide.shapes.title.text = "Unsupervised Learning Results"
table = slide.shapes.add_table(
    rows=4,
    cols=3,
    left=Inches(0.5),
    top=Inches(1.5),
    width=Inches(9),
    height=Inches(4.0),
).table

table.cell(0, 0).text = "Cluster"
table.cell(0, 1).text = "Records"
table.cell(0, 2).text = "Default Rate"

table.cell(1, 0).text = "low_risk"
table.cell(1, 1).text = "41,823"
table.cell(1, 2).text = "0.1199"

table.cell(2, 0).text = "medium_risk"
table.cell(2, 1).text = "29,702"
table.cell(2, 2).text = "0.1893"

table.cell(3, 0).text = "high_risk"
table.cell(3, 1).text = "48,475"
table.cell(3, 2).text = "0.2673"


# 10. Add generated chart slides if images exist.
chart_files = [
    ["Risk Distribution", "risk_distribution.png", "Borrower status distribution"],
    ["Cluster Visualization", "cluster_scatter.png", "K-Means borrower clusters"],
    ["Association Rules", "association_rules.png", "Top association rules"],
]

for chart in chart_files:
    title = chart[0]
    image_name = chart[1]
    caption = chart[2]

    slide = prs.slides.add_slide(prs.slide_layouts[5])
    slide.shapes.title.text = title
    image_path = visuals_dir / image_name

    if image_path.exists():
        slide.shapes.add_picture(str(image_path), Inches(0.7), Inches(1.5), width=Inches(8.6))

    text_box = slide.shapes.add_textbox(Inches(0.7), Inches(6.5), Inches(8.5), Inches(0.5))
    text_box.text_frame.text = caption


# 11. Add system implementation slide.
slide = prs.slides.add_slide(prs.slide_layouts[1])
slide.shapes.title.text = "System Implementation"
text_frame = slide.placeholders[1].text_frame
text_frame.clear()
bullets = [
    "Default probability prediction",
    "Stage classification",
    "Interest recommendation",
    "What-if simulation",
]
for index, bullet in enumerate(bullets):
    if index == 0:
        paragraph = text_frame.paragraphs[0]
    else:
        paragraph = text_frame.add_paragraph()
    paragraph.text = bullet
    paragraph.font.size = Pt(22)


# 12. Add architecture slide.
slide = prs.slides.add_slide(prs.slide_layouts[5])
slide.shapes.title.text = "Architecture Flow"
steps = [
    "Dataset",
    "Preprocessing",
    "ETL",
    "Warehouse",
    "Association Mining",
    "Classification",
    "Clustering",
    "Visualization",
    "System Output",
]

for index, step in enumerate(steps):
    column_number = index % 3
    row_number = index // 3
    left = 0.5 + column_number * 3.0
    top = 2.0 + row_number * 1.3

    shape = slide.shapes.add_shape(
        MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE,
        Inches(left),
        Inches(top),
        Inches(2.2),
        Inches(0.7),
    )
    shape.text_frame.text = step


# 13. Add conclusion slide.
slide = prs.slides.add_slide(prs.slide_layouts[1])
slide.shapes.title.text = "Conclusion"
text_frame = slide.placeholders[1].text_frame
text_frame.clear()
bullets = [
    "The project covers DET syllabus Units I-IV",
    "It includes preprocessing, ETL, association mining, classification, and clustering",
    "The final system demonstrates practical loan risk prediction",
]
for index, bullet in enumerate(bullets):
    if index == 0:
        paragraph = text_frame.paragraphs[0]
    else:
        paragraph = text_frame.add_paragraph()
    paragraph.text = bullet
    paragraph.font.size = Pt(22)


# 14. Save the presentation.
try:
    prs.save(str(output_file))
    print(f"Presentation created at: {output_file}")
except PermissionError:
    fallback_file = output_file.with_name("Loan_Risk_Project_Presentation_Updated.pptx")
    prs.save(str(fallback_file))
    print(f"Presentation created at: {fallback_file}")
